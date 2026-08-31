#!/usr/bin/env python3
"""Fill the Score column (H) of an official UCLouvain file from a custom file.

Usage:
    python3 fill_grades.py TARGET.xlsx CUSTOM.xlsx [-o OUTPUT.xlsx]
                           [--sheet NAME] [--email-col LETTER] [--score-col LETTER]

- TARGET : official-format file (e.g. session_2025_2_LEPL1503.xlsx),
           with the 'Email' header in column G and 'Score' in column H.
- CUSTOM : your own file holding the emails and the points.
           The email column is detected automatically; the points column too
           (you are prompted only if it is ambiguous).

The result is written to a NEW file (the target is never modified) and keeps a
single sheet -> no 500 error on upload.

The module is split so that the core logic is pure and testable: interactive
prompts are injected as `chooser` callbacks, defaulting to a stdin prompt.
"""
import argparse
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

# An email is a token with a single '@' and at least one dot in the domain part.
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# Non-numeric score codes allowed by the official file legend.
SCORE_CODES = {"A", "T", "S", "M"}


def norm_email(value):
    """Return the normalized email (lowercase, stripped) or None if not an email."""
    if value is None:
        return None
    s = str(value).strip()
    return s.lower() if EMAIL_RE.match(s) else None


def looks_like_score(value):
    """Return True if the value looks like a point/score (number or A/T/S/M code)."""
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    s = str(value).strip()
    if not s:
        return False
    if s.upper() in SCORE_CODES:
        return True
    try:
        float(s.replace(",", "."))
        return True
    except ValueError:
        return False


def prompt_chooser(prompt, options):
    """Default interactive chooser: list options and read a 1-based index from stdin."""
    for i, opt in enumerate(options, 1):
        print(f"  [{i}] {opt}")
    while True:
        raw = input(f"{prompt} ").strip()
        if raw.isdigit() and 1 <= int(raw) <= len(options):
            return int(raw) - 1
        print("  Invalid choice, try again.")


def pick_sheet(wb, forced=None, chooser=prompt_chooser):
    """Pick the worksheet to read from the custom workbook.

    If `forced` is given it must exist. With a single sheet it is used directly.
    Otherwise `chooser` decides among the sheet names.
    """
    if forced is not None:
        if forced not in wb.sheetnames:
            raise ValueError(f"sheet '{forced}' not found; sheets: {wb.sheetnames}")
        return wb[forced]
    if len(wb.sheetnames) == 1:
        return wb[wb.sheetnames[0]]
    idx = chooser("Which sheet to use?", wb.sheetnames)
    return wb[wb.sheetnames[idx]]


def detect_email_column(ws):
    """Return the 1-based index of the column holding the most emails, or None."""
    counts = {}
    for col in range(1, ws.max_column + 1):
        n = sum(1 for row in range(1, ws.max_row + 1)
                if norm_email(ws.cell(row, col).value))
        if n:
            counts[col] = n
    if not counts:
        return None
    return max(counts, key=lambda c: counts[c])


def detect_score_column(ws, email_col, chooser=prompt_chooser):
    """Detect the points column: the one (excluding emails) with the most score-like
    values. If several columns are close, `chooser` disambiguates."""
    scored = {}
    for col in range(1, ws.max_column + 1):
        if col == email_col:
            continue
        n = sum(1 for row in range(1, ws.max_row + 1)
                if looks_like_score(ws.cell(row, col).value))
        if n:
            scored[col] = n
    if not scored:
        return None
    best = max(scored.values())
    # Columns within 10% of the best are considered ambiguous candidates.
    leaders = sorted(c for c, n in scored.items() if n >= best * 0.9)
    if len(leaders) == 1:
        return leaders[0]
    labels = []
    for c in leaders:
        sample = [ws.cell(r, c).value for r in range(1, ws.max_row + 1)
                  if ws.cell(r, c).value not in (None, "")][:4]
        labels.append(f"column {get_column_letter(c)} — e.g. {sample}")
    idx = chooser("Which column holds the points?", labels)
    return leaders[idx]


def build_mapping(ws, email_col, score_col):
    """Build the email -> point mapping from a custom worksheet.

    Returns (mapping, duplicates) where `duplicates` lists emails seen more than
    once with conflicting values (the last value wins in the mapping).
    """
    mapping = {}
    duplicates = []
    for row in range(1, ws.max_row + 1):
        email = norm_email(ws.cell(row, email_col).value)
        if not email:
            continue
        point = ws.cell(row, score_col).value
        if email in mapping and mapping[email] != point:
            duplicates.append(email)
        mapping[email] = point
    return mapping, duplicates


def read_custom(path, forced_sheet=None, forced_email_col=None,
                forced_score_col=None, chooser=prompt_chooser):
    """Open the custom workbook and return (mapping, duplicates, info).

    `info` is a dict with the resolved sheet title and column indices, useful for
    logging and tests.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = pick_sheet(wb, forced_sheet, chooser)

    email_col = (column_index_from_string(forced_email_col) if forced_email_col
                 else detect_email_column(ws))
    if not email_col:
        raise ValueError(f"no email column detected in sheet '{ws.title}'")

    score_col = (column_index_from_string(forced_score_col) if forced_score_col
                 else detect_score_column(ws, email_col, chooser))
    if not score_col:
        raise ValueError("no points column detected in the custom file")

    mapping, duplicates = build_mapping(ws, email_col, score_col)
    info = {"sheet": ws.title, "email_col": email_col, "score_col": score_col}
    return mapping, duplicates, info


def find_data_start(ws, email_col):
    """Return the first row (1-based) that holds an email in `email_col`, or None."""
    for row in range(1, ws.max_row + 1):
        if norm_email(ws.cell(row, email_col).value):
            return row
    return None


def fill_target(ws, mapping, email_col, score_col):
    """Write points into `score_col` of the target worksheet by matching emails.

    Returns a report dict with the counts and the various warning lists.
    """
    start = find_data_start(ws, email_col)
    if start is None:
        raise ValueError(
            f"no email found in column {get_column_letter(email_col)} of the target")

    written = 0
    missing_in_custom = []   # target emails absent from the custom file
    decimal_warn = []        # points with a non-integer value
    used_emails = set()
    for row in range(start, ws.max_row + 1):
        email = norm_email(ws.cell(row, email_col).value)
        if not email:
            continue
        if email in mapping:
            point = mapping[email]
            ws.cell(row, score_col).value = point
            used_emails.add(email)
            written += 1
            if isinstance(point, float) and not point.is_integer():
                decimal_warn.append((email, point))
        else:
            missing_in_custom.append(email)

    missing_in_target = sorted(set(mapping) - used_emails)
    return {
        "written": written,
        "missing_in_custom": missing_in_custom,
        "missing_in_target": missing_in_target,
        "decimal_warn": decimal_warn,
    }


def run(target_path, custom_path, output_path=None, sheet=None,
        email_col=None, score_col=None, target_email_col="G",
        target_score_col="H", chooser=prompt_chooser):
    """End-to-end run: read custom, fill target, save output. Returns a report dict
    (also carrying `duplicates`, `custom_info` and `output`)."""
    target_path = Path(target_path).expanduser()
    custom_path = Path(custom_path).expanduser()
    out_path = (Path(output_path).expanduser() if output_path
                else target_path.with_name(target_path.stem + "_filled.xlsx"))

    if not target_path.exists():
        raise FileNotFoundError(f"target not found: {target_path}")
    if not custom_path.exists():
        raise FileNotFoundError(f"custom not found: {custom_path}")

    mapping, duplicates, info = read_custom(
        custom_path, sheet, email_col, score_col, chooser)

    twb = openpyxl.load_workbook(target_path)
    tws = twb.active
    if tws is None:
        raise ValueError("target workbook has no active sheet")

    report = fill_target(
        tws, mapping,
        column_index_from_string(target_email_col),
        column_index_from_string(target_score_col))

    twb.save(out_path)
    report["duplicates"] = sorted(set(duplicates))
    report["custom_info"] = info
    report["mapping"] = mapping
    report["output"] = out_path
    return report


def print_report(report, mapping=None):
    """Pretty-print the run report and its warnings."""
    print("\n" + "=" * 60)
    print(f"OK: {report['written']} points written -> {report['output']}")
    print("=" * 60)

    if report["missing_in_custom"]:
        print(f"\n⚠️  {len(report['missing_in_custom'])} student(s) in the target with NO "
              f"point in the custom file (Score left empty):")
        for e in report["missing_in_custom"]:
            print(f"    - {e}")
    else:
        print("\n✅  All emails in the target found a corresponding point "
              "in the custom file.")

    if report["missing_in_target"]:
        print(f"\n⚠️  {len(report['missing_in_target'])} email(s) in the custom file ABSENT "
              f"from the target (points ignored):")
        for e in report["missing_in_target"]:
            extra = f"  (point={mapping[e]!r})" if mapping else ""
            print(f"    - {e}{extra}")

    if report["duplicates"]:
        print(f"\n⚠️  {len(report['duplicates'])} duplicate email(s) in the custom file "
              f"(last value kept):")
        for e in report["duplicates"]:
            print(f"    - {e}")

    if report["decimal_warn"]:
        print(f"\n⚠️  {len(report['decimal_warn'])} decimal point(s) (decimals may be "
              f"rejected for this learning unit):")
        for e, pt in report["decimal_warn"]:
            print(f"    - {e} = {pt}")

    if not (report["missing_in_custom"] or report["missing_in_target"]
            or report["duplicates"] or report["decimal_warn"]):
        print("\n✅  No warning: all emails match on both sides.")


def main(argv=None):
    p = argparse.ArgumentParser(
        description="Fill the target Score column from the custom file.")
    p.add_argument("target", help="Official file (.xlsx)")
    p.add_argument("custom", help="Your custom file (.xlsx)")
    p.add_argument("-o", "--output", help="Output file (default: <target>_filled.xlsx)")
    p.add_argument("--sheet", help="Custom sheet name (asked if several and omitted)")
    p.add_argument("--email-col", help="Force the custom email column (e.g. B)")
    p.add_argument("--score-col", help="Force the custom points column (e.g. C)")
    p.add_argument("--target-email-col", default="G", help="Target email column (default G)")
    p.add_argument("--target-score-col", default="H", help="Target score column (default H)")
    args = p.parse_args(argv)

    try:
        report = run(
            args.target, args.custom, args.output, args.sheet,
            args.email_col, args.score_col,
            args.target_email_col, args.target_score_col)
    except (ValueError, FileNotFoundError) as exc:
        sys.exit(f"Error: {exc}")

    info = report["custom_info"]
    print(f"\nCustom: sheet '{info['sheet']}', emails in column "
          f"{get_column_letter(info['email_col'])}, points in column "
          f"{get_column_letter(info['score_col'])}.")
    print_report(report, report["mapping"])


if __name__ == "__main__":
    main()
