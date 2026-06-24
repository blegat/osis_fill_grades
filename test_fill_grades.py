"""Tests for fill_grades.

Workbooks are built on the fly in a temporary directory so no real student data
is needed. The official "target" layout is reproduced faithfully: metadata rows,
the table header on row 13 (Email in column G, Score in column H) and data from
row 14 onward.
"""
import openpyxl
import pytest

import fill_grades as fg


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def make_target(path, students):
    """Build a target workbook mimicking the official format.

    `students` is a list of (email, score) tuples; score is usually None (empty)
    since the whole point is to fill it.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    # A few metadata rows like the real file.
    ws["A1"] = "2025-26 - LEPL1503 Projet 3"
    ws["A2"] = "Session: 2"
    ws["A6"] = "Score"
    # Header row 13: only the columns we rely on need to be exact.
    ws["A13"] = "Academic year"
    ws["F13"] = "Name"
    ws["G13"] = "Email"
    ws["H13"] = "Score"
    for i, (email, score) in enumerate(students):
        row = 14 + i
        ws.cell(row, 7).value = email   # column G
        ws.cell(row, 8).value = score   # column H
    wb.save(path)
    return path


def make_custom(path, rows, header=("Name", "Email", "Points"), sheets=None):
    """Build a custom workbook. If `sheets` is given, it is a dict of
    name -> rows and produces a multi-sheet workbook instead."""
    wb = openpyxl.Workbook()
    if sheets is None:
        sheets = {"mes_points": rows}
    first = True
    for name, srows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        if header:
            ws.append(list(header))
        for r in srows:
            ws.append(list(r))
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# Pure helpers
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("value,expected", [
    ("john.doe@uclouvain.be", "john.doe@uclouvain.be"),
    ("  John.DOE@UCLouvain.be ", "john.doe@uclouvain.be"),
    ("not an email", None),
    ("missing@dot", None),
    ("@nope.be", None),
    (None, None),
    (42, None),
])
def test_norm_email(value, expected):
    assert fg.norm_email(value) == expected


@pytest.mark.parametrize("value,expected", [
    (18, True),
    (12.5, True),
    ("15", True),
    ("12,5", True),   # comma decimal
    ("A", True),      # absent
    ("t", True),      # cheating, case-insensitive
    ("", False),
    ("hello", False),
    (None, False),
])
def test_looks_like_score(value, expected):
    assert fg.looks_like_score(value) is expected


# --------------------------------------------------------------------------- #
# Column detection
# --------------------------------------------------------------------------- #
def test_detect_email_column(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", [
        ("Albert", "a@uclouvain.be", 10),
        ("Bauer", "b@uclouvain.be", 11),
    ])
    ws = openpyxl.load_workbook(path)["mes_points"]
    assert fg.detect_email_column(ws) == 2  # column B


def test_detect_email_column_none(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", [("x", "y", 1)], header=None)
    ws = openpyxl.load_workbook(path).active
    assert fg.detect_email_column(ws) is None


def test_detect_score_column_auto(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", [
        ("Albert", "a@uclouvain.be", 10),
        ("Bauer", "b@uclouvain.be", 11),
    ])
    ws = openpyxl.load_workbook(path)["mes_points"]
    assert fg.detect_score_column(ws, email_col=2) == 3  # column C


def test_detect_score_column_ambiguous_uses_chooser(tmp_path):
    # Two numeric columns of equal strength -> chooser must decide.
    path = make_custom(tmp_path / "c.xlsx",
                       [("a@uclouvain.be", 1, 2),
                        ("b@uclouvain.be", 3, 4)],
                       header=("Email", "Col1", "Col2"))
    ws = openpyxl.load_workbook(path).active
    picked = {}

    def chooser(prompt, options):
        picked["options"] = options
        return 1  # choose the second candidate

    col = fg.detect_score_column(ws, email_col=1, chooser=chooser)
    assert len(picked["options"]) == 2
    assert col == 3  # column C (the second numeric column)


# --------------------------------------------------------------------------- #
# Sheet selection
# --------------------------------------------------------------------------- #
def test_pick_sheet_single(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", [("a@uclouvain.be", 1)],
                       header=("Email", "Points"))
    wb = openpyxl.load_workbook(path)
    assert fg.pick_sheet(wb).title == "mes_points"


def test_pick_sheet_multiple_uses_chooser(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", None, header=("Email", "Points"),
                       sheets={"first": [("a@uclouvain.be", 1)],
                               "second": [("b@uclouvain.be", 2)]})
    wb = openpyxl.load_workbook(path)
    chosen = fg.pick_sheet(wb, chooser=lambda prompt, opts: opts.index("second"))
    assert chosen.title == "second"


def test_pick_sheet_forced_missing(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", [("a@uclouvain.be", 1)],
                       header=("Email", "Points"))
    wb = openpyxl.load_workbook(path)
    with pytest.raises(ValueError):
        fg.pick_sheet(wb, forced="nope")


# --------------------------------------------------------------------------- #
# Mapping
# --------------------------------------------------------------------------- #
def test_build_mapping_case_insensitive_and_duplicates(tmp_path):
    path = make_custom(tmp_path / "c.xlsx", [
        ("Albert", "alice@uclouvain.be", 10),
        ("Bauer", "BOB@uclouvain.be", 11),
        ("Dup", "alice@uclouvain.be", 99),   # conflicting duplicate
    ])
    ws = openpyxl.load_workbook(path)["mes_points"]
    mapping, dups = fg.build_mapping(ws, email_col=2, score_col=3)
    assert mapping["alice@uclouvain.be"] == 99   # last value wins
    assert mapping["bob@uclouvain.be"] == 11      # normalized key
    assert dups == ["alice@uclouvain.be"]


# --------------------------------------------------------------------------- #
# Filling the target
# --------------------------------------------------------------------------- #
def test_fill_target_writes_scores_and_warnings(tmp_path):
    target = make_target(tmp_path / "t.xlsx", [
        ("julian.albert@student.uclouvain.be", None),
        ("AMOS.ALLARD@student.uclouvain.be", None),   # different case
        ("augustin.barras@student.uclouvain.be", None),
        ("nobody@student.uclouvain.be", None),        # absent from custom
    ])
    mapping = {
        "julian.albert@student.uclouvain.be": 18,
        "amos.allard@student.uclouvain.be": 15,
        "augustin.barras@student.uclouvain.be": 12.5,   # decimal
        "ghost@student.uclouvain.be": 7,                # absent from target
    }
    ws = openpyxl.load_workbook(target).active
    report = fg.fill_target(ws, mapping, email_col=7, score_col=8)

    assert report["written"] == 3
    assert ws.cell(14, 8).value == 18
    assert ws.cell(15, 8).value == 15            # matched despite case
    assert ws.cell(16, 8).value == 12.5
    assert ws.cell(17, 8).value is None          # nobody -> left empty
    assert report["missing_in_custom"] == ["nobody@student.uclouvain.be"]
    assert report["missing_in_target"] == ["ghost@student.uclouvain.be"]
    assert report["decimal_warn"] == [("augustin.barras@student.uclouvain.be", 12.5)]


def test_fill_target_no_email_raises(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "empty.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "empty.xlsx").active
    with pytest.raises(ValueError):
        fg.fill_target(ws, {}, email_col=7, score_col=8)


# --------------------------------------------------------------------------- #
# End to end
# --------------------------------------------------------------------------- #
def test_run_end_to_end(tmp_path):
    target = make_target(tmp_path / "session.xlsx", [
        ("a@student.uclouvain.be", None),
        ("b@student.uclouvain.be", None),
    ])
    custom = make_custom(tmp_path / "custom.xlsx", [
        ("Alice", "a@student.uclouvain.be", 16),
        ("Bob", "b@student.uclouvain.be", 14),
    ])
    out = tmp_path / "out.xlsx"
    report = fg.run(target, custom, output_path=out)

    assert out.exists()
    wb = openpyxl.load_workbook(out)
    # Single sheet preserved -> no 500 on upload.
    assert wb.sheetnames == ["Sheet"]
    ws = wb.active
    assert ws.cell(14, 8).value == 16
    assert ws.cell(15, 8).value == 14
    assert report["written"] == 2
    assert report["custom_info"]["email_col"] == 2
    assert report["custom_info"]["score_col"] == 3
    assert report["output"] == out


def test_run_default_output_name(tmp_path):
    target = make_target(tmp_path / "session.xlsx",
                         [("a@student.uclouvain.be", None)])
    custom = make_custom(tmp_path / "custom.xlsx",
                         [("Alice", "a@student.uclouvain.be", 16)])
    report = fg.run(target, custom)
    assert report["output"] == tmp_path / "session_filled.xlsx"
    assert report["output"].exists()


def test_run_missing_target(tmp_path):
    custom = make_custom(tmp_path / "custom.xlsx",
                         [("Alice", "a@student.uclouvain.be", 16)])
    with pytest.raises(FileNotFoundError):
        fg.run(tmp_path / "does_not_exist.xlsx", custom)


def test_run_does_not_modify_target(tmp_path):
    target = make_target(tmp_path / "session.xlsx",
                         [("a@student.uclouvain.be", None)])
    custom = make_custom(tmp_path / "custom.xlsx",
                         [("Alice", "a@student.uclouvain.be", 16)])
    fg.run(target, custom, output_path=tmp_path / "out.xlsx")
    # The original target's Score cell must still be empty.
    ws = openpyxl.load_workbook(target).active
    assert ws.cell(14, 8).value is None
